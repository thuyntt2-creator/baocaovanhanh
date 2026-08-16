import docx
import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

docx_path = r'C:\Users\lap4all\Downloads\NTB - Báo Cáo Tuần (HRBP - ARD) - Tuần 32.docx'
excel_path = r'C:\Users\lap4all\Downloads\BaoCao_Tuan_NTB_W32_2026.xlsx'

doc = docx.Document(docx_path)
wb = openpyxl.load_workbook(excel_path, data_only=True)

print("=== CHECKING ALL TABLES FOR W31 / W32 HEADERS AND CONTENT ===")
for idx, table in enumerate(doc.tables):
    rows = len(table.rows)
    cols = len(table.columns)
    first_row_text = " | ".join([c.text.strip().replace('\n', ' ') for c in table.rows[0].cells])
    second_row_text = " | ".join([c.text.strip().replace('\n', ' ') for c in table.rows[1].cells]) if rows > 1 else ""
    
    has_w31 = "W31" in first_row_text or "W31" in second_row_text or "Tuần 31" in first_row_text or "Tuần 31" in second_row_text
    has_w32 = "W32" in first_row_text or "W32" in second_row_text or "Tuần 32" in first_row_text or "Tuần 32" in second_row_text
    
    print(f"Table {idx:2d} ({rows:2d}x{cols:2d}): W31={has_w31}, W32={has_w32} | R0: {first_row_text[:80]}")
